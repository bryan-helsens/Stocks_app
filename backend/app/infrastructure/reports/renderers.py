"""Renderers turning a ReportDocument into PDF / CSV / Excel / JSON bytes.

PDF uses WeasyPrint (HTML/CSS → PDF) and is imported lazily so the app starts
even when the native WeasyPrint libraries are absent; an informative error is
raised only when a PDF is actually requested.
"""

from __future__ import annotations

import csv
import io
import json
from html import escape

from app.core.errors import ExternalServiceError
from app.domain.value_objects.enums import ReportFormat
from app.infrastructure.reports.document import ReportDocument, Section, Table


def render(document: ReportDocument, fmt: ReportFormat | str) -> tuple[bytes, str]:
    """Render *document* to *fmt*; return (bytes, mime_type)."""
    f = str(fmt).upper()
    if f == ReportFormat.JSON:
        return _json(document), "application/json"
    if f == ReportFormat.CSV:
        return _csv(document), "text/csv"
    if f == ReportFormat.EXCEL:
        return _excel(document), (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    if f == ReportFormat.PDF:
        return _pdf(document), "application/pdf"
    raise ExternalServiceError(f"Unsupported report format: {fmt}")


# --------------------------------------------------------------------------- #
# JSON
# --------------------------------------------------------------------------- #
def _json(doc: ReportDocument) -> bytes:
    payload = {
        "title": doc.title,
        "subtitle": doc.subtitle,
        "generated_at": doc.generated_at.isoformat(),
        "disclaimer": doc.disclaimer,
        "sections": [
            {
                "heading": s.heading,
                "facts": s.facts,
                "text": s.text,
                "tables": [
                    {"title": t.title, "headers": t.headers, "rows": t.rows} for t in s.tables
                ],
            }
            for s in doc.sections
        ],
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


# --------------------------------------------------------------------------- #
# CSV (flattens facts and tables into a single sheet)
# --------------------------------------------------------------------------- #
def _csv(doc: ReportDocument) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([doc.title])
    if doc.subtitle:
        w.writerow([doc.subtitle])
    w.writerow(["Generated", doc.generated_at.isoformat()])
    w.writerow([])
    for section in doc.sections:
        w.writerow([f"# {section.heading}"])
        for key, value in section.facts.items():
            w.writerow([key, value])
        for table in section.tables:
            w.writerow([])
            w.writerow([table.title])
            w.writerow(table.headers)
            w.writerows(table.rows)
        w.writerow([])
    w.writerow(["Disclaimer", doc.disclaimer])
    return buf.getvalue().encode("utf-8")


# --------------------------------------------------------------------------- #
# Excel (one summary sheet + one sheet per table)
# --------------------------------------------------------------------------- #
def _excel(doc: ReportDocument) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = doc.title
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = doc.subtitle
    summary["A3"] = f"Generated: {doc.generated_at.isoformat()}"
    row = 5
    for section in doc.sections:
        summary.cell(row=row, column=1, value=section.heading).font = Font(bold=True)
        row += 1
        for key, value in section.facts.items():
            summary.cell(row=row, column=1, value=key)
            summary.cell(row=row, column=2, value=value)
            row += 1
        row += 1

    sheet_names: set[str] = {"Summary"}
    for section in doc.sections:
        for table in section.tables:
            ws = wb.create_sheet(_unique_sheet_name(table.title, sheet_names))
            ws.append(table.headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for r in table.rows:
                ws.append(r)

    summary.cell(row=row + 1, column=1, value="Disclaimer")
    summary.cell(row=row + 1, column=2, value=doc.disclaimer)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _unique_sheet_name(title: str, used: set[str]) -> str:
    base = (title or "Sheet")[:28].replace("/", "-").replace("\\", "-") or "Sheet"
    name = base
    i = 1
    while name in used:
        name = f"{base[:25]}_{i}"
        i += 1
    used.add(name)
    return name


# --------------------------------------------------------------------------- #
# PDF (HTML + WeasyPrint)
# --------------------------------------------------------------------------- #
def _pdf(doc: ReportDocument) -> bytes:
    html = _to_html(doc)
    try:
        from weasyprint import HTML  # lazy: native deps may be absent
    except Exception as exc:  # pragma: no cover - optional native dependency
        raise ExternalServiceError(
            "PDF generation is unavailable (WeasyPrint not installed)."
        ) from exc
    return HTML(string=html).write_pdf()


def _to_html(doc: ReportDocument) -> str:
    parts = [
        "<html><head><meta charset='utf-8'><style>",
        "body{font-family:Inter,Arial,sans-serif;color:#0F172A;margin:32px;}",
        "h1{font-size:22px;margin-bottom:0;} .sub{color:#64748B;margin-top:4px;}",
        "h2{font-size:15px;border-bottom:1px solid #E2E8F0;padding-bottom:4px;margin-top:24px;}",
        "table{border-collapse:collapse;width:100%;margin:8px 0;font-size:12px;}",
        "th,td{border:1px solid #E2E8F0;padding:6px 8px;text-align:left;}",
        "th{background:#F1F5F9;}",
        ".facts td:first-child{color:#64748B;width:40%;}",
        ".disclaimer{margin-top:28px;font-size:11px;color:#92400E;background:#FEF3C7;"
        "padding:10px 12px;border-radius:6px;}",
        "</style></head><body>",
        f"<h1>{escape(doc.title)}</h1>",
        f"<div class='sub'>{escape(doc.subtitle)} · {doc.generated_at:%Y-%m-%d %H:%M} UTC</div>",
    ]
    for section in doc.sections:
        parts.append(_section_html(section))
    parts.append(f"<div class='disclaimer'>⚠️ {escape(doc.disclaimer)}</div>")
    parts.append("</body></html>")
    return "".join(parts)


def _section_html(section: Section) -> str:
    parts = [f"<h2>{escape(section.heading)}</h2>"]
    if section.text:
        parts.append(f"<p>{escape(section.text)}</p>")
    if section.facts:
        parts.append("<table class='facts'>")
        for key, value in section.facts.items():
            parts.append(f"<tr><td>{escape(key)}</td><td>{escape(value)}</td></tr>")
        parts.append("</table>")
    for table in section.tables:
        parts.append(f"<h3 style='font-size:13px;'>{escape(table.title)}</h3>")
        parts.append("<table><tr>")
        parts.extend(f"<th>{escape(h)}</th>" for h in table.headers)
        parts.append("</tr>")
        for row in table.rows:
            parts.append("<tr>" + "".join(f"<td>{escape(str(c))}</td>" for c in row) + "</tr>")
        parts.append("</table>")
    return "".join(parts)


__all__ = ["render", "ReportDocument", "Section", "Table"]
